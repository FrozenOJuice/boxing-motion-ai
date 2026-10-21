"""Inspected source provenance only; all frame mappings remain provisional."""

from xml.etree import ElementTree
from zipfile import ZipFile

from openpyxl import load_workbook

from ..boxingvi import integer_cell
from .recording import v5_recording


def value_extent(path):
    """Ignore Excel's million-row formatting dimensions; inspect actual value/formula cells."""
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    with ZipFile(path) as archive:
        sheets = [
            n
            for n in archive.namelist()
            if n.startswith("xl/worksheets/") and n.endswith(".xml") and "/_rels/" not in n
        ]
        if len(sheets) != 1:
            raise ValueError("Expected one inspected worksheet part")
        root = ElementTree.fromstring(archive.read(sheets[0]))
    last, width = 1, 0
    for row in root.iter(ns + "row"):
        for cell in row.findall(ns + "c"):
            if any(
                "".join(child.itertext())
                for child in cell
                if child.tag in {ns + "v", ns + "f", ns + "is"}
            ):
                last = max(last, int(row.attrib["r"]))
                letters = "".join(c for c in cell.attrib["r"] if c.isalpha())
                col = 0
                for c in letters:
                    col = col * 26 + ord(c) - 64
                width = max(width, col)
    return last, width


def read_v5_annotations(path):
    return read_annotations(path, v5_recording())


def read_annotations(path, recording):
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        if book.sheetnames != ["Sheet1"]:
            raise ValueError(
                "Expected inspected recording layout: one Sheet1; review changed workbook"
            )
        last, width = value_extent(path)
        if width > len(recording.header):
            raise ValueError("Workbook width differs from inspected layout")
        rows, auxiliary_rows = [], []
        for n, cells in enumerate(
            book["Sheet1"].iter_rows(max_row=last, max_col=len(recording.header), values_only=True),
            1,
        ):
            if n == 1:
                if list(cells) != recording.header:
                    raise ValueError("Workbook header differs from explicit inspected layout")
                continue
            if not any(v is not None for v in cells):
                continue
            if len(cells) != len(recording.header):
                raise ValueError("Workbook width differs from inspected layout")
            source_cells = {chr(65 + i): value for i, value in enumerate(cells)}
            start, end, label = [
                source_cells[recording.columns[k]] for k in ("start", "end", "class")
            ]
            helpers = {k: v for k, v in source_cells.items() if k not in recording.columns.values()}
            if start is None and end is None and label is None:
                auxiliary_rows.append({"source_row": n, "source_cells": source_cells})
                continue
            helper_start, helper_end = helpers.get("D"), helpers.get("E")
            flags = ["provisional_mapping", "class_unverified", "boundaries_unverified"]
            start_index = end_index = None
            try:
                a, b = integer_cell(start), integer_cell(end)
                numeric = a >= 0 and b >= 0
                if numeric:
                    start_index, end_index = a + recording.frame_offset, b + recording.frame_offset
            except ValueError:
                numeric = False
            if not numeric:
                flags.append("invalid_frame_values")
            elif end < start:
                flags.append("reversed_bounds")
            elif end == start:
                flags.append("equal_bounds_end_convention_unverified")
            if not isinstance(label, str):
                flags.append("invalid_source_class")
            elif label not in recording.class_map and label != "Jab ":
                flags.append("unrecognized_source_class")
            if isinstance(label, str) and label != label.strip():
                flags.append("source_class_whitespace")
            if any(
                helpers.get(col) != formula.format(row=n)
                for col, formula in recording.helper_formulas.items()
            ):
                flags.append("unexpected_helper_formula")
            if any(
                source_cells[left] != source_cells[right]
                for left, right in recording.mirrored_columns.items()
            ):
                flags.append("mirrored_annotation_mismatch")
            rows.append(
                {
                    "source_row": n,
                    "raw_start": start,
                    "raw_end": end,
                    "source_class": label,
                    "candidate_start_index": start_index,
                    "candidate_end_index": end_index,
                    "helper_D": helper_start,
                    "helper_E": helper_end,
                    "source_helpers": helpers,
                    "source_cells": source_cells,
                    "flags": flags,
                }
            )
        valid = [r for r in rows if "invalid_frame_values" not in r["flags"]]
        previous = None
        seen = set()
        for row in valid:
            start, end = row["raw_start"], row["raw_end"]
            if previous is not None and start < previous:
                row["flags"].append("out_of_row_order")
            previous = start
            key = (start, end, row["source_class"])
            if key in seen:
                row["flags"].append("duplicate")
            seen.add(key)
            if end >= start:
                # Inclusive endpoint comparison is deliberately conservative for review only.
                overlap = [
                    other["source_row"]
                    for other in valid
                    if other is not row
                    and other["raw_end"] >= other["raw_start"]
                    and max(start, other["raw_start"]) <= min(end, other["raw_end"])
                ]
                if overlap:
                    row["flags"].append("suspect_overlap_inclusive_candidate")
                    row["overlapping_source_rows"] = overlap
        return {
            "schema_version": "recording-provisional-1.1",
            "video_id": recording.video_id,
            "recording_id": recording.video_id,
            "recording_config": recording.model_dump(),
            "sheet": "Sheet1",
            "mapping_status": "provisional_configured_for_review",
            "display_candidate": (
                f"raw {recording.columns['start']}/{recording.columns['end']} + "
                f"{recording.frame_offset} -> decoded indices; provisional"
            ),
            "verified_author_fps": None,
            "verified_frame_base": None,
            "verified_end_convention": None,
            "verified_trim_offset": None,
            "boxer_id": recording.boxer_id,
            "session_id": recording.session_id,
            "canonical_class_map": None,
            "unlabeled_intervals": "unknown; never automatically no_punch",
            "auxiliary_only_rows": auxiliary_rows,
            "rows": rows,
        }
    finally:
        book.close()
