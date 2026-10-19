"""Inspect XLSX without guessing a release schema; parse only explicit mappings."""

from collections import Counter
from pathlib import Path
from typing import Literal

from openpyxl import load_workbook
from pydantic import Field, field_validator

from .models import Contract, Event, PunchType, relative_path


class SheetMapping(Contract):
    workbook: str
    sheet: str = Field(min_length=1)
    data_start_row: int = Field(ge=1, strict=True)
    columns: dict[str, int]
    video: str
    boxer_id: str | None = Field(default=None, min_length=1)
    session_id: str | None = Field(default=None, min_length=1)
    frame_base: Literal[0, 1]
    end_convention: Literal["inclusive", "exclusive"]
    class_map: dict[str, PunchType] = Field(min_length=1)

    @field_validator("workbook", "video")
    @classmethod
    def paths(cls, value):
        return relative_path(value)

    @field_validator("columns")
    @classmethod
    def column_indices(cls, value):
        if set(value) != {"frame_start", "frame_end", "punch_type"}:
            raise ValueError("columns requires exactly frame_start, frame_end, punch_type")
        if any(type(v) is not int or not 1 <= v <= 16384 for v in value.values()):
            raise ValueError("Column positions must be integers in 1..16384")
        if len(set(value.values())) != 3:
            raise ValueError("Column positions must be distinct")
        return value


class Mapping(Contract):
    annotation_version: str = Field(min_length=1)
    evidence: str = Field(min_length=1)  # Operator's verification note, not author evidence.
    sheets: list[SheetMapping] = Field(min_length=1)


def integer_cell(value) -> int:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise ValueError(f"Frame must be an integer numeric cell, got {value!r}")
    if not float(value).is_integer():
        raise ValueError(f"Nonintegral frame: {value!r}")
    return int(value)


def inspect_and_parse(path: Path, root: Path, mapping: Mapping | None, issue):
    """Return sheet inventory, canonical events, and raw mapped class counts."""
    rel = path.relative_to(root).as_posix()
    inventory, events, classes = [], [], Counter()
    configs = [s for s in mapping.sheets if s.workbook == rel] if mapping else []
    book = None
    try:
        book = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        for config in configs:
            if config.sheet not in book.sheetnames:
                issue("missing_sheet", "error", rel, f"Sheet {config.sheet!r} is absent")
        for sheet in book:
            choices = [c for c in configs if c.sheet == sheet.title]
            if len(choices) > 1:
                issue("duplicate_mapping", "error", rel, f"Multiple mappings: {sheet.title}")
            config = choices[0] if len(choices) == 1 else None
            if config is None:
                issue(
                    "schema_needed",
                    "error",
                    f"{rel}:{sheet.title}",
                    "Inspect preview and supply --mapping with sheet, row/column positions, "
                    "exact class map, media path and verified frame conventions",
                )
            rows, preview = 0, []
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                if not any(v is not None for v in row):
                    continue
                rows += 1
                if len(preview) < 5:
                    preview.append(
                        {
                            "row": row_number,
                            "cells": [str(v)[:200] if v is not None else None for v in row[:20]],
                        }
                    )
                if config is None or row_number < config.data_start_row:
                    continue
                location = f"{rel}:{sheet.title}:{row_number}"
                try:
                    values = {
                        key: row[index - 1] if index <= len(row) else None
                        for key, index in config.columns.items()
                    }
                    label = str(values["punch_type"])
                    classes[label] += 1
                    if label not in config.class_map:
                        issue(
                            "unrecognized_class",
                            "error",
                            location,
                            f"Unmapped class {label!r}; do not infer handedness or no_punch",
                        )
                        continue
                    events.append(
                        Event(
                            event_id=location,
                            video_id=config.video,
                            boxer_id=config.boxer_id,
                            session_id=config.session_id,
                            frame_start=integer_cell(values["frame_start"]),
                            frame_end=integer_cell(values["frame_end"]),
                            frame_base=config.frame_base,
                            end_convention=config.end_convention,
                            punch_type=config.class_map[label],
                            annotation_version=mapping.annotation_version,
                            source_workbook=rel,
                            source_sheet=sheet.title,
                            source_row=row_number,
                            source_class=label,
                        )
                    )
                except (ValueError, TypeError, OverflowError) as exc:
                    issue("malformed_row", "error", location, str(exc))
            inventory.append(
                {
                    "workbook": rel,
                    "sheet": sheet.title,
                    "nonempty_rows": rows,
                    "preview": preview,
                    "mapped": config is not None,
                }
            )
    except Exception as exc:  # Bad ZIP/XML/encrypted/unsupported workbooks must produce a report.
        issue("unreadable_workbook", "error", rel, f"{type(exc).__name__}: {exc}")
    finally:
        if book is not None:
            book.close()
    return inventory, events, classes
