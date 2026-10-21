"""Synthetic second-recording transfer; no real data or model is needed."""

import json
from fractions import Fraction
from types import SimpleNamespace

import pytest
from openpyxl import Workbook
from typer.testing import CliRunner

from boxing_motion_ai.cli import app
from boxing_motion_ai.data.pose_quality import POLICIES, assess_quality
from boxing_motion_ai.data.preparation import prepare_recording
from boxing_motion_ai.vision.annotations import read_annotations
from boxing_motion_ai.vision.contracts import DecodedFrame
from boxing_motion_ai.vision.pipeline import extract_recording, sha256
from boxing_motion_ai.vision.recording import SOURCE_CLASSES, RecordingConfig
from boxing_motion_ai.vision.sampling import plan_sample


@pytest.fixture
def v4(tmp_path):
    source = tmp_path / "source"
    source.mkdir()
    video, workbook, model = [source / n for n in ["V4.mkv", "V4.xlsx", "test.task"]]
    video.write_bytes(b"synthetic, not video")
    model.write_bytes(b"synthetic, not model")
    classes = {("Lead hook" if k == "Lead Hook" else k): v for k, v in SOURCE_CLASSES.items()}
    config = RecordingConfig(
        video_id="V4",
        header=["Start", "End", "Class", "Difference", "Start_norm", "End_norm"],
        helper_formulas={"D": "=B{row}-A{row}", "E": "=A{row}-6437", "F": "=B{row}-6437"},
        class_map=classes,
        frame_offset=0,
        end_assumption="inclusive",
    )
    book = Workbook()
    sheet = book.active
    sheet.title = "Sheet1"
    sheet.append(config.header)
    for i, label in enumerate(classes):
        n = i + 2
        sheet.append([i * 5 + 2, i * 5 + 3, label, f"=B{n}-A{n}", f"=A{n}-6437", f"=B{n}-6437"])
    sheet.append([40, 41, "Jab", "=B8-A8", "=A8-6437", "=B8-6437"])
    book.save(workbook)
    return video, workbook, model, config


def decoder(path, width):
    for i in range(50):
        yield DecodedFrame(
            i,
            round(i * 1001 / 24),
            Fraction(1, 1000),
            640,
            360,
            SimpleNamespace(shape=(360, 640, 3)),
        )


class Backend:
    def __init__(self, *args, **kwargs):
        pass

    def detect(self, rgb, ms):
        points = [(0.5, 0.5, 0.1, 0.99, 0.99)] * 33
        points[13] = (0.5, 0.5, 0.1, 0.1, 0.99)
        points[15] = (0.5, 0.5, 0.1, 0.1, 0.99)
        return points

    def close(self):
        pass


def test_sample_transfer_preserves_sparse_indices_masks_and_unknown_identity(v4, tmp_path):
    np = pytest.importorskip("numpy")
    video, workbook, model, recording = v4
    original = [sha256(p) for p in (video, workbook, model)]
    run, prepared, report = [tmp_path / n for n in ["run", "prepared", "quality"]]
    ranges = [(i * 5 + 2, i * 5 + 4) for i in range(6)]
    extract_recording(
        video,
        workbook,
        model,
        run,
        ranges=ranges,
        recording=recording,
        decoder=decoder,
        backend_factory=Backend,
        overlay_limit=0,
    )
    with pytest.raises(ValueError, match="full-video"):
        prepare_recording(run, workbook, prepared)
    result = prepare_recording(run, workbook, prepared, allow_sample=True)
    assert result["source_frame_count"] is None and result["extracted_frame_count"] == 12
    assert result["boxer_id"] is None and result["recording_id"] == "V4"
    rows = [json.loads(s) for s in (prepared / "candidates.jsonl").read_text().splitlines()]
    assert rows[0]["source_interval_half_open"] == [2, 4]
    assert rows[0]["array_slice_half_open"] == [0, 2]
    assert rows[0]["recording_id"] == "V4" and rows[0]["boxer_id"] is None
    assert rows[-1]["annotation_eligible"] and not rows[-1]["quality_assessable"]
    assert rows[-1]["pose_quality"] is None  # unextracted, NOT missing
    with np.load(prepared / "poses.npz") as arrays:
        assert arrays["frame_index"].tolist() == [n for a, b in ranges for n in range(a, b)]
        assert not arrays["valid_mask"][:, 13].any()
        assert (arrays["confidence"][:, 13, 0] == 0.1).all()
    quality = assess_quality(prepared, report, 2)
    assert quality["policies"] == POLICIES
    assert all(c["retention"]["mask_aware70"] == 1 for c in quality["by_class"].values())
    assert quality["visual_review_selection"][0]["frames"] == [2, 3]
    assert [sha256(p) for p in (video, workbook, model)] == original


def test_v4_layout_helpers_and_frame_offset_explicit(v4):
    _, workbook, _, config = v4
    data = read_annotations(workbook, config)
    assert data["rows"][0]["source_helpers"]["F"] == "=B2-6437"
    assert "unexpected_helper_formula" not in data["rows"][0]["flags"]
    assert data["rows"][2]["source_class"] == "Lead hook"
    shifted = config.model_copy(update={"frame_offset": -1})
    assert read_annotations(workbook, shifted)["rows"][0]["candidate_start_index"] == 1
    with pytest.raises(ValueError, match="header"):
        read_annotations(workbook, config.model_copy(update={"header": [None] * 6}))


def test_sample_plan_deterministic_and_safe(v4, tmp_path):
    _, workbook, _, config = v4
    first = plan_sample(workbook, config, tmp_path / "a", 2)
    second = plan_sample(workbook, config, tmp_path / "b", 2)
    assert first == second and len(first["selected_rows"]) == 7
    assert json.loads((tmp_path / "a" / "pose-config.json").read_text())["confidence"] == 0.5
    with pytest.raises(ValueError, match="new plan"):
        plan_sample(workbook, config, tmp_path / "a")


def test_other_video_requires_explicit_recording(v4, tmp_path):
    video, workbook, model, _ = v4
    result = CliRunner().invoke(
        app,
        [
            "extract-poses",
            "--video",
            str(video),
            "--workbook",
            str(workbook),
            "--model",
            str(model),
            "--output",
            str(tmp_path / "out"),
        ],
    )
    assert result.exit_code == 1 and "--recording is required" in result.output


def test_generic_full_run_has_own_identity_and_cut_inventory(v4, tmp_path):
    video, workbook, model, recording = v4
    run, prepared = tmp_path / "full", tmp_path / "prepared-full"
    result = extract_recording(
        video,
        workbook,
        model,
        run,
        recording=recording,
        decoder=decoder,
        backend_factory=Backend,
        overlay_limit=0,
        scene_cuts=[42],
    )
    assert result["source_frame_count"] == 50
    assert result["config"]["scene_cuts"] == [42]
    report = prepare_recording(run, workbook, prepared)
    assert report["source_frame_count"] == 50 and report["video_id"] == "V4"
    assert report["confirmed_scene_cuts"] == [42]
    assert report["class_counts"]["jab"]["quality_assessable"] == 1  # row8 near cut
    assert "Assessed" in (prepared / "preparation.md").read_text()


def test_alternate_columns_auxiliary_rows_and_excel_formatting_extent(tmp_path):
    config = RecordingConfig(
        video_id="V1",
        header=["Start", None, "End", None, "Class", None],
        columns={"start": "A", "end": "C", "class": "E"},
        helper_formulas={"F": "=C{row}-A{row}"},
        class_map={"cross": "cross"},
        frame_offset=0,
        end_assumption="inclusive",
    )
    path = tmp_path / "V1.xlsx"
    book = Workbook()
    s = book.active
    s.title = "Sheet1"
    s.append(config.header)
    s.append([10, None, 13, None, "cross", "=C2-A2"])
    s.append([None, None, None, None, None, 123])
    s.cell(1048576, 6).number_format = "0.00"  # formatting is not a million annotations
    book.save(path)
    rows = read_annotations(path, config)
    assert len(rows["rows"]) == 1 and len(rows["auxiliary_only_rows"]) == 1
    assert rows["rows"][0]["raw_end"] == 13 and rows["rows"][0]["source_class"] == "cross"
    assert rows["rows"][0]["source_cells"]["E"] == "cross"


def test_mirrored_columns_are_checked_without_correction(tmp_path):
    config = RecordingConfig(
        video_id="V2",
        header=[None] * 5,
        helper_formulas={},
        class_map={"Jab": "jab"},
        mirrored_columns={"D": "A", "E": "B"},
        frame_offset=0,
        end_assumption="inclusive",
    )
    book = Workbook()
    s = book.active
    s.title = "Sheet1"
    s.append(config.header)
    s.append([10, 14, "Jab", 11, 14])
    path = tmp_path / "V2.xlsx"
    book.save(path)
    row = read_annotations(path, config)["rows"][0]
    assert "mirrored_annotation_mismatch" in row["flags"]
    assert row["raw_start"] == 10 and row["source_cells"]["D"] == 11


def test_flagged_class_sample_is_diagnostic_only(v4, tmp_path):
    _, workbook, _, config = v4
    b = __import__("openpyxl").load_workbook(workbook)
    b.active.cell(3, 1).value = 2
    b.active.cell(3, 2).value = 3  # overlaps row2
    b.save(workbook)
    result = plan_sample(workbook, config, tmp_path / "plan", 2)
    flagged = [r for r in result["selected_rows"] if r.get("selected_for_review_only")]
    assert {r["source_class"] for r in flagged} == {"Cross"}
    assert all("suspect_overlap_inclusive_candidate" in r["flags"] for r in flagged)
